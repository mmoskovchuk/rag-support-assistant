import pytest
from docx import Document as WordDocument

from support_rag import extraction
from support_rag.extraction import NO_PAGE, ExtractionError, extract_pages


def make_pdf(path, page_texts):
    """Write a minimal valid PDF; an empty string produces a page without a text layer (like a scan)."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in page_texts:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET" if text else ""
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        content_ref = len(objects)
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_ref} 0 R >>"
        )
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"

    body, offsets = b"%PDF-1.4\n", []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(body))
        body += f"{number} 0 obj\n{obj}\nendobj\n".encode()
    xref = len(body)
    body += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    body += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    body += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(body)


def test_digital_pdf_uses_text_layer_and_ocr_only_for_scanned_pages(tmp_path, settings, monkeypatch):
    ocr_calls = []

    def fake_ocr(path, number, settings):
        ocr_calls.append(number)
        return "Recognised text from the scanned page."

    monkeypatch.setattr(extraction, "ocr_pdf_page", fake_ocr)
    pdf = tmp_path / "mixed.pdf"
    make_pdf(pdf, ["Annual leave is twenty four calendar days.", ""])

    pages = extract_pages(pdf, settings)

    assert [(p.page_number, p.method) for p in pages] == [(1, "text"), (2, "ocr")]
    assert "twenty four" in pages[0].text
    assert ocr_calls == [2]  # the page with a text layer never touched OCR


def test_docx_keeps_paragraphs_and_tables_in_order(tmp_path, settings):
    path = tmp_path / "policy.docx"
    word = WordDocument()
    word.add_heading("Відпустки", level=1)
    word.add_paragraph("Щорічна відпустка становить 24 дні.")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Стаж", "Днів"
    table.cell(1, 0).text, table.cell(1, 1).text = "5 років", "26"
    word.add_paragraph("Кінець документа.")
    word.save(path)

    [page] = extract_pages(path, settings)

    assert page.page_number == NO_PAGE
    assert page.text.index("24 дні") < page.text.index("Стаж | Днів") < page.text.index("Кінець")
    assert "5 років | 26" in page.text


def test_broken_word_file_raises_extraction_error(tmp_path, settings):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip archive")
    with pytest.raises(ExtractionError):
        extract_pages(path, settings)
